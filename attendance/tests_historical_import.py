from datetime import date, timedelta
from io import StringIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

import openpyxl
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from django.utils import timezone
from rest_framework.test import APIClient

from accounts.models import User
from applications.models import Application
from attendance.models import Attendance, AbsenceWarning
from attendance.services.student_scope import attendance_metrics, student_attendance_queryset
from attendance.tasks import finalize_meet_attendance_task
from cohorts.models import Cohort
from courses.models import Course


class HistoricalImportTests(TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.seed = Path(self.tmp.name)/"seed.xlsx"
        self.register = Path(self.tmp.name)/"register.xlsx"
        self.course = Course.objects.create(code="VLSI-DESIGN", name="VLSI", status="PUBLISHED")
        self.cohort = Cohort.objects.create(code="G2-26", course=self.course, start_date=date(2026,1,1), end_date=date(2027,1,1), status="TRAINING")
        self.students=[]
        for email in ["alice@example.com", "bob@example.com", "new@example.com"]:
            u=User.objects.create_user(email=email, role="STUDENT", password="Pass@12345!")
            p=u.student_profile
            Application.objects.create(student=p, course=self.course, assigned_cohort=self.cohort, status="TRAINING")
            self.students.append(p)
        self.mapping=patch("attendance.management.commands.import_historical_attendance.LegacySeed.EXPLICIT_STUDENT_MAP", {"alice@example.com":"ALICE", "bob@example.com":"BOB"})
        self.mapping.start()
        self.addCleanup(self.mapping.stop)
        wb=openpyxl.Workbook();ws=wb.active;ws.append(["Email"]);ws.append(["alice@example.com"]);ws.append(["bob@example.com"]);wb.save(self.seed);wb.close()
        wb=openpyxl.Workbook();ws=wb.active;ws.title="Attendance Register"
        ws.append(["Historical attendance"]);ws.append(["No.","Student Name","Cohort","One","Two"])
        ws.append([1,"ALICE","G2-26","P","A"]);ws.append([2,"BOB","G2-26","A","P"])
        daily=wb.create_sheet("Daily Summary");daily.append(["#","Date","Session"])
        self.dates=[timezone.localdate()-timedelta(days=30),timezone.localdate()-timedelta(days=29)]
        daily.append([1,self.dates[0],"One"]);daily.append([2,self.dates[1],"Two"]);wb.save(self.register);wb.close()

    def run_import(self, **extra):
        out=StringIO()
        call_command("import_historical_attendance", attendance_file=str(self.register), seed_file=str(self.seed), cohort_id=str(self.cohort.pk), expected_students=2, expected_sessions=2, stdout=out, **extra)
        return json.loads(out.getvalue())

    def apply(self):
        plan=self.run_import()
        return self.run_import(apply=True, confirm_plan=plan["plan_sha256"])

    def test_dry_run_and_repeated_import_preserve_history_without_side_effects(self):
        passwords=list(User.objects.order_by("id").values_list("password",flat=True))
        plan=self.run_import()
        self.assertEqual(Attendance.objects.count(),0)
        self.assertEqual(plan["marks"],4)
        self.apply()
        again=self.run_import(apply=True, confirm_plan=plan["plan_sha256"])
        self.assertEqual(again["existing"],2)
        self.assertEqual(Attendance.objects.count(),2)
        self.assertFalse(AbsenceWarning.objects.exists())
        self.assertEqual(passwords,list(User.objects.order_by("id").values_list("password",flat=True)))
        for row in Attendance.objects.all():
            self.assertIsNone(row.google_meet_attendance_data)
            self.assertFalse(row.meeting_link)
            self.assertFalse(row.joined_students.exists())

    def test_present_and_absent_imports_visible_only_to_mapped_students(self):
        self.apply()
        for p in self.students[:2]:
            self.assertEqual(student_attendance_queryset(p).count(),2)
            self.assertEqual(attendance_metrics(p,p.applications.first()),{"total":2,"present":1,"percentage":50.0,"arithmetic_percentage":50.0})
        Application.objects.filter(student=self.students[2]).update(applied_at=timezone.now()-timedelta(days=100))
        self.students[2].student_identity_issued_at=None
        self.students[2].save(update_fields=["student_identity_issued_at"])
        self.assertEqual(student_attendance_queryset(self.students[2]).count(),0)
        from django.contrib.admin.sites import AdminSite
        from students.admin import StudentProfileAdmin
        from students.models import StudentProfile
        self.assertIn("50.0% (1/2)", StudentProfileAdmin(StudentProfile, AdminSite()).attendance_rate_display(self.students[0]))
        client=APIClient();client.force_authenticate(self.students[0].user)
        result=client.get("/api/attendance/").data["results"]
        self.assertEqual(len(result),2)
        for row in result:
            self.assertNotIn("historical_attendance_data",row)
            self.assertNotIn("google_meet_attendance_data",row)
            self.assertIsNone(row["start_time"])
            self.assertEqual(row["student_dashboard_data"]["source"],"HISTORICAL_REGISTER")
            self.assertNotIn(str(self.students[1].pk),json.dumps(row,default=str))

    def test_live_sessions_and_staff_assignments_are_not_changed(self):
        live=Attendance.objects.create(cohort=self.cohort,course=self.course,title="Live",class_date=timezone.localdate(),start_time="20:00",meeting_link="https://meet.google.com/abc-defg-hij")
        before=list(Attendance.objects.filter(pk=live.pk).values())
        self.apply()
        self.assertEqual(before,list(Attendance.objects.filter(pk=live.pk).values()))
        self.assertFalse(self.cohort.mentors.exists())

    def test_overlapping_session_or_wrong_plan_blocks_entire_import(self):
        plan=self.run_import()
        with self.assertRaises(CommandError):self.run_import(apply=True,confirm_plan="incorrect")
        Attendance.objects.create(cohort=self.cohort,title="Existing",class_date=self.dates[1],start_time="20:00")
        with self.assertRaises(CommandError):self.run_import(apply=True,confirm_plan=plan["plan_sha256"])
        self.assertEqual(Attendance.objects.count(),1)

    def test_blank_mark_never_becomes_an_absence(self):
        wb=openpyxl.load_workbook(self.register);wb["Attendance Register"]["D3"]=None;wb.save(self.register);wb.close()
        with self.assertRaises(CommandError):self.apply()
        self.assertEqual(Attendance.objects.count(),0)

    def test_google_finalizer_does_not_overwrite_history_or_send_warnings(self):
        self.apply()
        with patch("attendance.services.real_meet_attendance_service.RealMeetAttendanceService.get_structured_attendance") as google:
            result=finalize_meet_attendance_task(str(Attendance.objects.first().pk))
        google.assert_not_called()
        self.assertEqual(result["status"],"SKIPPED_HISTORICAL")

